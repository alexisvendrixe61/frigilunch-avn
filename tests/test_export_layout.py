"""
The Excel layout the floor actually reads — the four changes asked for in the
04/09 meeting ("aanpassingen Lay-out om planning te kunnen lezen"), plus the
day break asked for after the first run shipped.

These build two tiny days and export them for real, so they check the written
workbook rather than the intent: the layout is the deliverable here, and it had
no test at all.
"""
from __future__ import annotations

from dataclasses import replace
from datetime import date, time

import openpyxl
import pytest

from src.exporter import export_multi_to_excel
from src.models import DaySchedule, PlannedOrder
from tests.conftest import make_order, make_product

DAG1 = date(2026, 8, 31)   # a Monday
DAG2 = date(2026, 9, 1)

DATUM_COL, OMSCHRIJVING_COL, LIJN_COL = 2, 4, 9
ALLERGIE_COL, HULP_MAT_COL = 14, 19


def planned(art_nr, *, lijn=1, start=(6, 0), eind=(8, 0), plan_nr="45225", **product_kwargs):
    product = make_product(art_nr, **product_kwargs)
    order = replace(make_order(art_nr, 2000, DAG1), plan_nr=plan_nr)
    return PlannedOrder(
        order=order, product=product, lijn=lijn,
        start=time(*start), eind=time(*eind), ombouw_voor=0, man=4.0,
    )


@pytest.fixture
def workbook(tmp_path):
    """Two days, two lines, one of them a sate production."""
    sate = planned("SATE", lijn=1, start=(14, 0), eind=(21, 0),
                   plan_nr="45061", allergie_niveau=6)
    dag1 = DaySchedule(
        datum=DAG1,
        lijn1=[planned("A", lijn=1, plan_nr="45225"), sate],
        lijn2=[planned("B", lijn=2, plan_nr="45041")],
    )
    dag2 = DaySchedule(
        datum=DAG2,
        lijn1=[planned("C", lijn=1, plan_nr="43645")],
        lijn2=[],
    )
    path = tmp_path / "planning.xlsx"
    export_multi_to_excel([dag1, dag2], str(path))
    return openpyxl.load_workbook(path)


def _continu(wb):
    return wb["planning maaltijden"]


def test_a_blank_row_separates_the_days(workbook):
    """~480 rows only read as a plan if the day blocks are visibly separated."""
    ws = _continu(workbook)
    dagen = [
        (r, ws.cell(row=r, column=DATUM_COL).value)
        for r in range(2, ws.max_row + 1)
    ]
    laatste_dag1 = max(r for r, d in dagen if d == "31/08/2026")
    eerste_dag2 = min(r for r, d in dagen if d == "01/09/2026")
    assert eerste_dag2 == laatste_dag1 + 2, "exactly one row between the two days"
    tussenrij = laatste_dag1 + 1
    assert all(
        ws.cell(row=tussenrij, column=c).value in (None, "")
        for c in range(1, ws.max_column + 1)
    ), "the separating row must be empty"


def test_the_number_column_holds_the_plan_number(workbook):
    ws = _continu(workbook)
    assert ws.cell(row=1, column=1).value == "plan nr."
    nrs = [
        ws.cell(row=r, column=1).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=3).value
    ]
    # Numbers, not text — the column has to sort and filter like the source
    # planning's does.
    assert nrs == [45225, 45061, 45041, 43645]


def test_hulp_mat_lists_the_assembly_devices(workbook, tmp_path):
    """Abbreviations from "machines assemblage", counts from the product db."""
    product = replace(
        make_product("DEV"),
        assemblage_toestellen={
            "multipond_l1_nodig": True, "leonard_nodig": True,
            "turbo_aantal": 2, "kaasstrooier_nodig": True,
        },
    )
    po = replace(planned("DEV"), product=product)
    path = tmp_path / "devices.xlsx"
    export_multi_to_excel([DaySchedule(datum=DAG1, lijn1=[po])], str(path))
    ws = openpyxl.load_workbook(path)["planning maaltijden"]

    assert ws.cell(row=1, column=HULP_MAT_COL).value == "hulp mat."
    assert ws.cell(row=2, column=HULP_MAT_COL).value == "I-L-2T-K"


def test_a_product_without_device_data_gets_a_blank_hulp_mat(workbook):
    """505 of the source planning's rows have none. Blank, never guessed.

    (openpyxl reads an empty cell back as None, whichever of the two was
    written.)"""
    ws = _continu(workbook)
    assert ws.cell(row=2, column=HULP_MAT_COL).value in (None, "")


def test_order_rows_carry_no_colour(workbook):
    """"kleuren weg" — the allergen level reads in its own column instead."""
    ws = _continu(workbook)
    for r in range(2, ws.max_row + 1):
        if not ws.cell(row=r, column=3).value:
            continue            # blank separator / sate marker
        cell = ws.cell(row=r, column=OMSCHRIJVING_COL)
        assert cell.fill.patternType is None, f"row {r} is filled"


def test_sate_gets_a_marker_row_like_the_source_planning(workbook):
    ws = _continu(workbook)
    markers = [
        r for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=OMSCHRIJVING_COL).value == "Saté"
    ]
    assert len(markers) == 1
    marker = markers[0]
    assert ws.cell(row=marker, column=OMSCHRIJVING_COL).font.b
    # ... directly above the sate production it marks.
    assert ws.cell(row=marker + 1, column=ALLERGIE_COL).value == "X"
