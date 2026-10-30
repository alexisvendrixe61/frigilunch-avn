"""
excel_export.py — Write planning output to Excel.

Public API:
    export_multi_to_excel(schedules, output_path, ...) -> None   native, no src bridge
    write_excel(outcome, output_dir, week) -> None               simplified stub
    write_action_csv(outcome, output_dir, week) -> None
"""
from __future__ import annotations

import csv
from pathlib import Path

from frigilunch.application import totals as _totals
from frigilunch.domain.constraints.allergen import is_sate
from frigilunch.domain.models import DaySchedule, PlannedOrder, SolveOutcome
from frigilunch.domain.state import order_identity

# ── colour palette (openpyxl hex, no #) ──────────────────────────────────────
_ALLERGIE_LABEL = {0: "-", 1: "!", 2: "!!", 3: "!!!", 4: "!!!!", 5: "!!!!!", 6: "X"}
_LIJN1_FILL  = "BDD7EE"
_LIJN2_FILL  = "E2CFEF"
_WARNING_FILL = "FFE0B2"
_SATE_LABEL  = "Saté"
_SATE_FILL   = "FF6699"

_NL_DAG = {0: "ma", 1: "di", 2: "wo", 3: "do", 4: "vr", 5: "za", 6: "zo"}

# ── column specs ──────────────────────────────────────────────────────────────
_HEADERS = [
    "plan nr.", "datum", "art.nr.", "omschrijving", "aantal", "lever datum",
    "CSB", "man", "lijn", "start hh:mm", "eind hh:mm", "st/hr",
    "turnover", "!!!!!", "V-A", "ma x st", "# we st in cha",
    "schaal", "hulp mat.", "kap vorm", "type karton", "#krt", "opmerking",
    "opmerking van tabblad orderverwerking", "Opmerking uit de analyselijst",
]

_COL_WIDTHS = {
    "A": 9, "B": 12, "C": 13, "D": 42, "E": 8, "F": 13,
    "G": 7, "H": 7, "I": 6, "J": 11, "K": 11, "L": 8,
    "M": 11, "N": 7, "O": 7, "P": 10, "Q": 14, "R": 14,
    "S": 12, "T": 13, "U": 14, "V": 6, "W": 40, "X": 40, "Y": 40,
}

_CONTINU_HEADERS = [
    "plan nr.", "datum", "art.nr.", "omschrijving", "aantal", "lever datum",
    "CSB", "man", "lijn", "start hh:mm", "eind hh:mm", "",
    "st/hr", "! ! ! ! !", "V-A", "max st/min", "# weegschalen",
    "schaal", "hulp mat.", "kap vorm", "type karton", "#krt",
    "opmerking van tabblad orderverwerking", "Opmerking uit de analyselijst",
    "uren steamer",
]

# ── assembly device "hulp mat." column ───────────────────────────────────────
# Order fixed to match the most frequent spellings in the source planning.
_HULP_MAT_VOLGORDE = [
    ("I",   ("multipond_l1_nodig", "isihda_l2_nodig")),
    ("L",   ("leonard_nodig",)),
    ("M",   ("multifill_nodig",)),
    ("H",   ("handtmann_aantal",)),
    ("T",   ("turbo_aantal",)),
    ("G",   ("grote_slicer_nodig",)),
    ("K",   ("kaasstrooier_nodig",)),
    ("CFS", ("cfs_pomp_nodig",)),
]


def _hulp_materiaal(product) -> str:
    counts = product.assemblage_toestellen or {}
    parts: list[str] = []
    for afkorting, keys in _HULP_MAT_VOLGORDE:
        n = max((int(counts.get(k) or 0) for k in keys), default=0)
        if n > 0:
            parts.append(f"{n}{afkorting}" if n > 1 else afkorting)
    return "-".join(parts)


def _plan_nr(po: PlannedOrder) -> int | str:
    nr = po.order.plan_nr
    return int(nr) if nr.isdigit() else nr


def _is_sate_order(po: PlannedOrder) -> bool:
    return is_sate(po.product)


# ── continuous planning sheet ─────────────────────────────────────────────────

def _write_continuous_planning_sheet(ws, schedules: list[DaySchedule]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    def mf(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    ws.append(_CONTINU_HEADERS)
    for cell in ws[1]:
        cell.fill = mf(_LIJN1_FILL)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = ["9", "12", "13", "42", "8", "13", "7", "7", "6", "11", "11",
              "6", "8", "9", "6", "10", "12", "13", "12", "12", "14", "6",
              "34", "34", "10"]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = int(w)

    rows: list[tuple] = []
    for s in schedules:
        for po in list(s.lijn1) + list(s.lijn2):
            rows.append((s.datum, po.lijn, po.start, s, po))
    rows.sort(key=lambda r: (r[0], r[1], r[2]))

    vorige_sate: tuple | None = None
    vorige_datum = None
    for datum, lijn, start, s, po in rows:
        allergie_label = _ALLERGIE_LABEL.get(po.product.allergie_niveau, "")
        krt = (po.order.aantal // po.product.stuks_per_karton
               if po.product.stuks_per_karton else 0)

        if vorige_datum is not None and datum != vorige_datum:
            ws.append([""] * len(_CONTINU_HEADERS))
        vorige_datum = datum

        if _is_sate_order(po):
            if vorige_sate != (datum, lijn):
                ws.append(["", "", "", _SATE_LABEL] + [""] * (len(_CONTINU_HEADERS) - 4))
                for cell in ws[ws.max_row]:
                    cell.fill = mf(_SATE_FILL)
                    cell.font = Font(bold=True)
            vorige_sate = (datum, lijn)
        else:
            vorige_sate = None

        ws.append([
            _plan_nr(po),
            s.datum.strftime("%d/%m/%Y"),
            po.order.art_nr,
            po.product.naam,
            po.order.aantal,
            po.order.lever_datum.strftime("%d/%m/%Y"),
            "",                             # CSB
            f"{po.man:.1f}",
            po.lijn,
            po.start.strftime("%H:%M"),
            po.eind.strftime("%H:%M"),
            po.ombouw_voor if po.ombouw_voor > 0 else "",
            po.product.snelheid,
            allergie_label,
            "", "", "",
            po.product.schaal,
            _hulp_materiaal(po.product),
            po.product.kapvorm,
            po.product.type_karton,
            krt,
            po.order.opmerking,
            po.order.opmerking_analyse,
            "",                             # uren steamer
        ])

    ws.freeze_panes = "A2"


# ── per-day sheet ─────────────────────────────────────────────────────────────

def _write_day_to_sheet(ws, schedule: DaySchedule) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    def mf(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    for col_letter, width in _COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    ws.append(_HEADERS)
    for cell in ws[ws.max_row]:
        cell.fill = mf(_LIJN1_FILL)
        cell.font = bold
        cell.alignment = center

    def write_lijn_header(lijn_nr):
        color = _LIJN1_FILL if lijn_nr == 1 else _LIJN2_FILL
        ws.append([f"LIJN {lijn_nr}"] + [""] * (len(_HEADERS) - 1))
        for cell in ws[ws.max_row]:
            cell.fill = mf(color)
            cell.font = bold

    def write_sate_label():
        ws.append(["", "", "", _SATE_LABEL] + [""] * (len(_HEADERS) - 4))
        for cell in ws[ws.max_row]:
            cell.fill = mf(_SATE_FILL)
            cell.font = bold

    def write_order_row(po: PlannedOrder):
        allergie_label = _ALLERGIE_LABEL.get(po.product.allergie_niveau, "")
        krt = (po.order.aantal // po.product.stuks_per_karton
               if po.product.stuks_per_karton else 0)
        opmerking = po.product.opmerking
        if po.warnings:
            suffix = " | ".join(po.warnings)
            opmerking = f"{opmerking} | {suffix}" if opmerking else suffix
        ws.append([
            _plan_nr(po),
            schedule.datum.strftime("%d/%m/%Y"),
            po.order.art_nr, po.product.naam, po.order.aantal,
            po.order.lever_datum.strftime("%d/%m/%Y"),
            "", f"{po.man:.1f}", po.lijn,
            po.start.strftime("%H:%M"), po.eind.strftime("%H:%M"),
            po.product.snelheid,
            po.ombouw_voor if po.ombouw_voor > 0 else "",
            allergie_label,
            "", "", "",
            po.product.schaal,
            _hulp_materiaal(po.product),
            po.product.kapvorm,
            po.product.type_karton, krt, opmerking,
            po.order.opmerking, po.order.opmerking_analyse,
        ])

    def write_lijn(lijn_nr: int, planned: list[PlannedOrder]):
        write_lijn_header(lijn_nr)
        vorige_sate = False
        for po in planned:
            if _is_sate_order(po):
                if not vorige_sate:
                    write_sate_label()
                vorige_sate = True
            else:
                vorige_sate = False
            write_order_row(po)

    write_lijn(1, schedule.lijn1)
    ws.append([""] * len(_HEADERS))
    write_lijn(2, schedule.lijn2)

    if schedule.warnings:
        ws.append([""] * len(_HEADERS))
        ws.append(["WARNINGS / MELDINGEN"] + [""] * (len(_HEADERS) - 1))
        for cell in ws[ws.max_row]:
            cell.fill = mf(_WARNING_FILL)
            cell.font = bold
        for w in schedule.warnings:
            ws.append(["", w] + [""] * (len(_HEADERS) - 2))

    kpis = schedule.kpis
    ws.append([""] * len(_HEADERS))
    ws.append([
        "KPI",
        f"L1: {kpis.get('lijn1_orders', 0)} orders "
        f"({kpis.get('lijn1_bezetting_pct', 0):.1f}% bezet)",
        f"L2: {kpis.get('lijn2_orders', 0)} orders "
        f"({kpis.get('lijn2_bezetting_pct', 0):.1f}% bezet)",
    ] + [""] * (len(_HEADERS) - 3))


# ── main export ───────────────────────────────────────────────────────────────

def export_multi_to_excel(
    schedules: list[DaySchedule],
    output_path: str,
    assignment_warnings: list[str] | None = None,
    product_warnings: list[str] | None = None,
    rejected_orders: list | None = None,
    rejection_reasons: dict | None = None,
) -> None:
    """Export a list of DaySchedule objects to a multi-tab Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as e:
        raise ImportError("openpyxl is vereist: pip install openpyxl") from e

    def mf(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    wb = Workbook()
    ws_ov = wb.active
    ws_ov.title = "Overzicht"

    # Order fulfillment summary
    planned_chunks = [po.order for s in schedules for po in (s.lijn1 + s.lijn2)]
    t = _totals.compute(
        planned_chunks + list(rejected_orders or []),
        planned_chunks,
        list(rejected_orders or []),
    )

    ws_ov.append(["ORDER FULFILLMENT SUMMARY"])
    for cell in ws_ov[ws_ov.max_row]:
        cell.fill = mf("E8F5E9")
        cell.font = Font(bold=True, size=12)

    ws_ov.append([
        f"Orders: {t.scheduled_orders + t.rejected_orders} | "
        f"Ingepland: {t.scheduled_orders} ({t.fulfillment_pct:.1f}%) | "
        f"Afgewezen: {t.rejected_orders}"
    ])
    ws_ov.append([
        f"Stuks: {t.scheduled_units + t.rejected_units:,} | "
        f"Ingepland: {t.scheduled_units:,} ({t.units_pct:.1f}%) | "
        f"Afgewezen: {t.rejected_units:,}"
    ])
    ws_ov.append([
        f"Productieruns: {t.scheduled_chunks} "
        f"({t.orders_split} order(s) gesplitst over meerdere dagen)"
    ])
    ws_ov.append([])

    ov_headers = [
        "Datum", "Dag", "Week", "L1 orders", "L1 %", "L2 orders", "L2 %",
        "Errors", "Warnings", "MEP items",
    ]
    ws_ov.append(ov_headers)
    for cell in ws_ov[ws_ov.max_row]:
        cell.fill = mf(_LIJN1_FILL)
        cell.font = Font(bold=True)

    for s in schedules:
        weekdag = _NL_DAG.get(s.datum.weekday(), "")
        week_nr = s.datum.isocalendar()[1]
        errors = sum(1 for w in s.warnings if w.startswith("[ERROR]"))
        warns  = sum(1 for w in s.warnings if "[WARN" in w)
        mep_n  = len(s.kpis.get("mep_items", []))
        ws_ov.append([
            s.datum.strftime("%d/%m/%Y"), weekdag, week_nr,
            s.kpis.get("lijn1_orders", 0),
            f"{s.kpis.get('lijn1_bezetting_pct', 0):.1f}%",
            s.kpis.get("lijn2_orders", 0),
            f"{s.kpis.get('lijn2_bezetting_pct', 0):.1f}%",
            errors, warns, mep_n,
        ])
        row = ws_ov.max_row
        if errors:
            for cell in ws_ov[row]:
                cell.fill = mf("FFC7CE")
        elif (s.kpis.get("lijn1_bezetting_pct", 0) > 90
              or s.kpis.get("lijn2_bezetting_pct", 0) > 90):
            for cell in ws_ov[row]:
                cell.fill = mf("FFCC99")

    if assignment_warnings:
        ws_ov.append([])
        ws_ov.append(["ONPLANBARE / OVERLADEN ORDERS"] + [""] * (len(ov_headers) - 1))
        for cell in ws_ov[ws_ov.max_row]:
            cell.fill = mf("FFC7CE")
            cell.font = Font(bold=True)
        for w in assignment_warnings:
            ws_ov.append(["", w])

    if product_warnings:
        ws_ov.append([])
        ws_ov.append(["AUTO-AANGEVULDE PRODUCTEN (check allergeenprofiel!)"]
                     + [""] * (len(ov_headers) - 1))
        for cell in ws_ov[ws_ov.max_row]:
            cell.fill = mf(_WARNING_FILL)
            cell.font = Font(bold=True)
        for w in product_warnings:
            ws_ov.append(["", w])

    ws_ov.column_dimensions["A"].width = 13
    ws_ov.column_dimensions["B"].width = 6
    ws_ov.column_dimensions["C"].width = 7
    for col in "DEFGHIJ":
        ws_ov.column_dimensions[col].width = 12

    # Continuous planning tab
    ws_continu = wb.create_sheet(title="planning maaltijden")
    _write_continuous_planning_sheet(ws_continu, schedules)

    # Per-day tabs
    for s in schedules:
        dag_label = _NL_DAG.get(s.datum.weekday(), "")
        sheet_name = f"{s.datum.strftime('%d-%m')} {dag_label}"
        ws_day = wb.create_sheet(title=sheet_name)
        ws_day.append([
            f"Planning  {dag_label.upper()}  {s.datum.strftime('%d/%m/%Y')}",
        ])
        for cell in ws_day[1]:
            cell.font = Font(bold=True, size=12)
            cell.fill = mf(_LIJN1_FILL if s.datum.weekday() < 4 else "E8F5E9")
        ws_day.append([""] * len(_HEADERS))
        _write_day_to_sheet(ws_day, s)

    # Rejected orders tab
    if rejected_orders:
        ws_rej = wb.create_sheet(title="Afgewezen Orders")
        rej_headers = ["plan nr.", "Art.Nr", "Klant", "Aantal", "Lever Datum",
                       "Besteld op", "Reden"]
        ws_rej.append(rej_headers)
        for cell in ws_rej[1]:
            cell.fill = mf("FFC7CE")
            cell.font = Font(bold=True, color="FFFFFF")

        if rejection_reasons is None:
            rejection_reasons = {}
        for order in rejected_orders:
            key = order_identity(order)
            reason = rejection_reasons.get(key, "Onbekend")
            ws_rej.append([
                int(order.plan_nr) if order.plan_nr.isdigit() else order.plan_nr,
                order.art_nr, order.klant, order.aantal,
                order.lever_datum.isoformat() if order.lever_datum else "",
                order.order_nr, reason,
            ])
        for col, w in zip("ABCDEFG", [9, 15, 20, 10, 15, 20, 50]):
            ws_rej.column_dimensions[col].width = w

    wb.save(output_path)
    print(f"Planning opgeslagen: {output_path}  "
          f"({len(schedules)} dagen, {wb.sheetnames[0]} + {len(schedules)} dag-tabs)")


# ── simplified write_excel (for testing / no full schedule) ──────────────────

_SIMPLE_HEADERS = [
    "plan nr.", "datum", "art.nr.", "omschrijving", "aantal",
    "lever datum", "lijn", "start", "eind", "allergen", "opmerking",
]
_ACTION_HEADERS = ["art.nr.", "klant", "aantal", "lever datum", "reden"]


def _po_row(po: PlannedOrder, datum: object) -> list:
    o, p = po.order, po.product
    return [
        o.plan_nr, str(datum), o.art_nr, p.naam, o.aantal,
        str(o.lever_datum), po.lijn,
        po.start.strftime("%H:%M"), po.eind.strftime("%H:%M"),
        _ALLERGIE_LABEL.get(p.allergie_niveau, "-"),
        o.opmerking,
    ]


def _write_day_sheet_simple(ws: object, day: DaySchedule) -> None:
    ws.append(_SIMPLE_HEADERS)
    for po in day.lijn1:
        ws.append(_po_row(po, day.datum))
    for po in day.lijn2:
        ws.append(_po_row(po, day.datum))
    if day.warnings:
        ws.append([])
        for w in day.warnings:
            ws.append(["WARN", w])


def write_excel(outcome: SolveOutcome, output_dir: Path, week: int) -> None:
    """Write one sheet per production day to planning_wkNN.xlsx (simplified)."""
    try:
        import openpyxl
    except ImportError:
        print("[WARN] openpyxl not installed — Excel export skipped")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    days: list[DaySchedule] = outcome.stats.get("days", [])
    if not days:
        ws = wb.create_sheet("(geen data)")
        ws.append(["Geen productiedagen beschikbaar"])
    else:
        for day in sorted(days, key=lambda d: d.datum):
            ws = wb.create_sheet(day.datum.strftime("%Y-%m-%d"))
            _write_day_sheet_simple(ws, day)

    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"planning_wk{week:02d}.xlsx"
    wb.save(path)
    print(f"[OK] Excel geschreven: {path}")


def write_action_csv(outcome: SolveOutcome, output_dir: Path, week: int) -> None:
    """Write rejected orders to actie_vereist_wkNN.csv."""
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"actie_vereist_wk{week:02d}.csv"
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(_ACTION_HEADERS)
        for order in outcome.rejected:
            writer.writerow([
                order.art_nr, order.klant, order.aantal,
                str(order.lever_datum), "rejected",
            ])
    print(f"[OK] Actie CSV geschreven: {path}")
